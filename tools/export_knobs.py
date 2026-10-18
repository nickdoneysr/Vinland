"""Export every input on the plan's Assumptions sheet to data/knobs.json.

Usage: python3 tools/export_knobs.py docs/Vinland_Game_Development_Plan.xlsx
The game reads tuning values from data/knobs.json, never from code literals (Backlog F01, Risk R2).
"""
import json, re, sys, pathlib, datetime
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent

def key(label):
    k = re.sub(r'\(.*?\)', '', str(label)).lower()
    k = re.sub(r"[^a-z0-9]+", '_', k).strip('_')
    return k

def main(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True)['Assumptions']
    out = {}
    for r in range(5, ws.max_row + 1):
        lab, val = ws.cell(r, 1).value, ws.cell(r, 2).value
        if lab is None or val is None or str(lab).isupper() or lab in ('Item', 'Knob', 'Phase'):
            continue
        if isinstance(val, datetime.datetime):
            val = val.date().isoformat()
        out[key(lab)] = val
    (ROOT / 'data' / 'knobs.json').write_text(json.dumps(out, indent=1, ensure_ascii=False))
    print(len(out), 'knobs')

if __name__ == '__main__':
    main(sys.argv[1])
